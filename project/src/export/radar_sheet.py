"""Build the attention-radar sheet of the delivery file from the signal_daily table and the user watchlist."""

from __future__ import annotations

import sqlite3
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

SHEET_NAME = "Radar chú ý"
LOOKBACK_DAYS = 6
MAX_ROWS = 300

# (key in signal_daily, display label, column width, number format)
RADAR_FIELDS: list[tuple[str, str, int, str | None]] = [
    ("trade_date", "Phiên", 12, None),
    ("entity_id", "Thực thể", 26, None),
    ("n_articles", "Số bài", 9, "0"),
    ("n_sources", "Số nguồn", 10, "0"),
    ("n_stories", "Số câu chuyện", 14, "0"),
    ("share", "Tỷ trọng", 10, "0.0%"),
    ("ama_z", "Chú ý bất thường (z)", 20, "0.0"),
    ("stale_ratio", "Tin đăng lại", 12, "0%"),
    ("n_updates", "Bản cập nhật", 13, "0"),
    ("net_sent_story", "Sentiment theo câu chuyện", 24, "0.00"),
    ("net_sent_volume", "Sentiment theo số bài", 22, "0.00"),
    ("net_sent_norm", "Sentiment hiệu chỉnh nguồn", 26, "0.00"),
    ("dispersion", "Phân tán giữa nguồn", 19, "0.00"),
    ("sent_shift", "Dịch chuyển sentiment", 20, "0.00"),
    ("first_source", "Nguồn đưa tin đầu", 24, None),
    ("coverage_pct", "Độ phủ phân tích (%)", 20, "0.0"),
]


def _next_session(d: date) -> date:
    """Return the next weekday after a date.

    Args:
        d: Reference date.

    Returns:
        Next date that falls on Monday to Friday.
    """
    nxt = d + timedelta(days=1)
    while nxt.weekday() >= 5:
        nxt += timedelta(days=1)
    return nxt


def radar_rows(conn: sqlite3.Connection, entity_ids: set[str], day: str) -> list[dict]:
    """Read the signal rows of the entities a user watches around the delivery date.

    A missing table or no rows returns an empty list and the sheet is skipped.

    Args:
        conn: Read-only connection to the operational database.
        entity_ids: Entity ids in the subscription of the user.
        day: Delivery date as ISO `YYYY-MM-DD`.

    Returns:
        At most 300 rows, newest first, and within a session the highest abnormal attention first.
    """
    if not entity_ids:
        return []
    try:
        d = date.fromisoformat(day)
    except ValueError:
        return []
    lo, hi = (d - timedelta(days=LOOKBACK_DAYS)).isoformat(), _next_session(d).isoformat()
    ids = sorted(entity_ids)
    rows: list[dict] = []
    try:
        for i in range(0, len(ids), 500):
            chunk = ids[i:i + 500]
            marks = ",".join("?" * len(chunk))
            cur = conn.execute(
                f"SELECT * FROM signal_daily WHERE trade_date BETWEEN ? AND ? "
                f"AND entity_id IN ({marks})", [lo, hi, *chunk])
            names = [c[0] for c in cur.description]
            rows += [dict(zip(names, r)) for r in cur.fetchall()]
    except sqlite3.OperationalError:
        return []
    rows.sort(key=lambda r: (r["trade_date"], r["ama_z"] if r["ama_z"] is not None else -1e9),
              reverse=True)
    return rows[:MAX_ROWS]


def add_radar_sheet(wb: Workbook, rows: list[dict]) -> None:
    """Add the attention-radar sheet to the workbook; do nothing when there are no rows.

    Args:
        wb: Delivery workbook that already has the main sheet.
        rows: Result of `radar_rows`.
    """
    if not rows:
        return
    ws = wb.create_sheet(SHEET_NAME)
    header_font, border = Font(bold=True), Border(bottom=Side(style="thin"))
    for col, (_k, label, width, _fmt) in enumerate(RADAR_FIELDS, start=1):
        cell = ws.cell(row=1, column=col, value=label)
        cell.font, cell.border = header_font, border
        cell.alignment = Alignment(vertical="center", horizontal="left")
        ws.column_dimensions[cell.column_letter].width = width
    for r_i, row in enumerate(rows, start=2):
        for col, (key, _label, _w, fmt) in enumerate(RADAR_FIELDS, start=1):
            val = row.get(key)
            cell = ws.cell(row=r_i, column=col)
            if isinstance(val, str):
                cell.value = val
                cell.data_type = "s"
            else:
                cell.value = val
                if fmt and val is not None:
                    cell.number_format = fmt
            cell.alignment = Alignment(vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(RADAR_FIELDS)).column_letter}{ws.max_row}"
