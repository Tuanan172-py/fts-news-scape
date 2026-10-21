"""Export the delivery report as a monochrome Excel (.xlsx) file."""
from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

from src.core.staging import safe_atomic_write
from src.export.radar_sheet import add_radar_sheet

SHEET_NAME = "Watchlist News"

# DELIVERY contract: (key in the row dict, display label, column width, wrap text).
# Block order: location, classification (short, for scanning and filtering), source, long text, technical.
DELIVERY_FIELDS: list[tuple[str, str, int, bool]] = [
    ("date",             "Date",               16, False),
    ("matched_entities", "Matched Entities",   20, False),
    # Three columns that expose the recognition source, placed next to the entity column
    # where readers look. `Intent — LLM` is what the model recognized from meaning,
    # `Intent — Code` is what the hard dictionary matched. What appears only in the first
    # column is the value the model adds: sub-brands, industries inferred from context, and
    # macro themes without keywords.
    ("intent_llm",       "Intent — LLM",       28, True),
    ("intent_code",      "Intent — Code",      22, True),
    ("intent_source",    "Intent Source",      18, False),
    ("title",            "Title",              48, True),
    ("sentiment",        "Sentiment",          12, False),
    ("time_sensitivity", "Time Sensitivity",   16, False),
    ("gold_status",      "Gold Status",        14, False),
    ("source_domain",    "Source",             16, False),
    ("summary",          "Summary",            64, True),
    ("key_points",       "Key Points",         64, True),
    ("implication",      "Market Implication", 46, True),
    ("url",              "URL",                38, False),
    ("article_id",       "Article ID",         20, False),
]

# Enum → English label. Raw / unexpected values are preserved as-is.
SENTIMENT_EN = {"positive": "Positive", "negative": "Negative", "neutral": "Neutral"}
TIME_SENSITIVITY_EN = {
    "urgent": "Urgent", "today": "Today", "this_week": "This Week",
    "this_month": "This Month", "archive": "Archive",
}
GOLD_STATUS_EN = {"GOLD": "Full", "L1_ONLY": "Preliminary"}

_EN_MAPS = {
    "sentiment": SENTIMENT_EN,
    "time_sensitivity": TIME_SENSITIVITY_EN,
    "gold_status": GOLD_STATUS_EN,
}
_VN_MAPS = _EN_MAPS  # Backward compatibility alias

_HEADER_FONT = Font(bold=True)
_HEADER_BORDER = Border(bottom=Side(style="thin"))
_HEADER_ALIGN = Alignment(vertical="center", horizontal="left")
_WRAP_ALIGN = Alignment(vertical="top", wrap_text=True)
_PLAIN_ALIGN = Alignment(vertical="top")


def en_label(field: str, value) -> str:
    """Convert a technical enum value to its English display label.

    Args:
        field: Field name to convert.
        value: Internal value to look up.

    Returns:
        Matching display label, or the original value as a string if it is not mapped.
    """
    m = _EN_MAPS.get(field)
    if not m:
        return "" if value is None else str(value)
    return m.get(value, "" if value in (None, "") else str(value))


vn_label = en_label  # Backward compatibility alias


def _set_text(cell, value: str, *, align) -> None:
    """Set a plain text value in a cell to prevent formula injection.

    Args:
        cell: Worksheet cell.
        value: Text to store.
        align: Alignment applied to the cell.
    """
    cell.value = value
    cell.data_type = "s"
    cell.alignment = align


def _parse_date(s: str) -> datetime | date | None:
    """Parse a date or date-time string.

    Args:
        s: String in the form 'YYYY-MM-DD HH:MM' or 'YYYY-MM-DD'.

    Returns:
        A datetime when the time part is present, a date when only the date is present, otherwise None.
    """
    raw = (s or "").strip()
    if len(raw) >= 16:
        try:
            return datetime.strptime(raw[:16], "%Y-%m-%d %H:%M")
        except (TypeError, ValueError):
            pass
    try:
        return datetime.strptime(raw[:10], "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return None


def build_workbook(rows: list[dict]) -> Workbook:
    """Build the monochrome Excel workbook from the aggregated articles.

    Args:
        rows: Article data dicts keyed by DELIVERY_FIELDS.

    Returns:
        Formatted openpyxl Workbook.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    for col, (_key, label, width, _wrap) in enumerate(DELIVERY_FIELDS, start=1):
        cell = ws.cell(row=1, column=col)
        _set_text(cell, label, align=_HEADER_ALIGN)
        cell.font = _HEADER_FONT
        cell.border = _HEADER_BORDER
        ws.column_dimensions[cell.column_letter].width = width

    for r_i, row in enumerate(rows, start=2):
        for col, (key, _label, _width, wrap) in enumerate(DELIVERY_FIELDS, start=1):
            cell = ws.cell(row=r_i, column=col)
            align = _WRAP_ALIGN if wrap else _PLAIN_ALIGN
            raw = row.get(key, "")

            if key == "date":
                d = _parse_date(raw)
                if d is not None:
                    cell.value = d
                    cell.number_format = "yyyy-mm-dd hh:mm" if isinstance(d, datetime) else "yyyy-mm-dd"
                    cell.alignment = align
                else:
                    _set_text(cell, str(raw or ""), align=align)
                continue

            if key in _EN_MAPS:
                _set_text(cell, en_label(key, raw), align=align)
                continue

            _set_text(cell, "" if raw is None else str(raw), align=align)
            # Clickable link without the Hyperlink style, to stay monochrome.
            if key == "url" and str(raw or "").startswith(("http://", "https://")):
                cell.hyperlink = str(raw)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(DELIVERY_FIELDS)).column_letter}{max(ws.max_row, 1)}"
    return wb


def write_delivery_xlsx(path: str | Path, rows: list[dict],
                        radar_rows: list[dict] | None = None) -> tuple[Path, bool]:
    """Write the table to an Excel (.xlsx) file through a safe atomic write.

    Args:
        path: Target file path.
        rows: Article data dicts.
        radar_rows: Signal rows for the attention-radar sheet; empty or None adds no sheet.

    Returns:
        Tuple of the path actually written and the flag that the write fell back to a snapshot because of a lock.
    """
    wb = build_workbook(rows)
    if radar_rows:
        add_radar_sheet(wb, radar_rows)
    return safe_atomic_write(Path(path), wb.save, binary=True, fallback_on_lock=True)
