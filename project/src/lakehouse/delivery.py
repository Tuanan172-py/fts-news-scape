"""Động cơ phân phối báo cáo Excel 2 sheets cá nhân hóa theo Watchlist người dùng."""

from __future__ import annotations

import concurrent.futures
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
import pyarrow.parquet as pq
import yaml

from src.core.staging import safe_atomic_write

SHEET_MAIN_NAME = "Tin tức trọng yếu"
SHEET_CITATIONS_NAME = "Luận điểm trích dẫn"

# Danh mục 15 cột giao hàng chuẩn mực cho Sheet 1
SHEET1_COLUMNS: list[tuple[str, str, int, bool]] = [
    ("date", "Date", 16, False),
    ("matched_entities", "Matched Entities", 20, False),
    ("intent_llm", "Intent — LLM", 28, True),
    ("intent_code", "Intent — Code", 22, True),
    ("intent_source", "Intent Source", 18, False),
    ("title", "Title", 48, True),
    ("sentiment", "Sentiment", 14, False),
    ("time_sensitivity", "Time Sensitivity", 16, False),
    ("gold_status", "Gold Status", 14, False),
    ("source_domain", "Source", 16, False),
    ("summary", "Summary", 64, True),
    ("key_points", "Key Points", 64, True),
    ("implication", "Market Implication", 46, True),
    ("url", "URL", 38, False),
    ("article_id", "Article ID", 22, False),
]

# Danh mục cột chi tiết cho Sheet 2 (Luận điểm & Trích dẫn nguyên văn)
SHEET2_COLUMNS: list[tuple[str, str, int, bool]] = [
    ("article_id", "Article ID", 22, False),
    ("date", "Date", 16, False),
    ("title", "Title", 44, True),
    ("matched_entities", "Matched Entities", 20, False),
    ("key_point", "Key Point / Luận điểm", 50, True),
    ("grounded_citation", "Grounded Citation / Trích dẫn nguyên văn", 60, True),
    ("source_domain", "Source", 16, False),
    ("url", "URL", 38, False),
]

# Bảng ánh xạ giá trị enum sang nhãn hiển thị doanh nghiệp
SENTIMENT_LABELS = {
    "positive": "Positive",
    "pos": "Positive",
    "negative": "Negative",
    "neg": "Negative",
    "neutral": "Neutral",
    "neu": "Neutral",
}

TIME_SENSITIVITY_LABELS = {
    "urgent": "Urgent",
    "urg": "Urgent",
    "today": "Today",
    "this_week": "This Week",
    "week": "This Week",
    "this_month": "This Month",
    "month": "This Month",
    "archive": "Archive",
    "arch": "Archive",
}

GOLD_STATUS_LABELS = {
    "GOLD": "Full",
    "L1_ONLY": "Preliminary",
}

_HEADER_FONT = Font(bold=True)
_HEADER_BORDER = Border(bottom=Side(style="thin"))
_HEADER_ALIGN = Alignment(vertical="center", horizontal="left")
_WRAP_ALIGN = Alignment(vertical="top", wrap_text=True)
_PLAIN_ALIGN = Alignment(vertical="top")


def _format_cell_text(cell: Any, value: str, align: Alignment) -> None:
    """Gán giá trị văn bản thuần vào ô để ngăn ngừa tiêm nhiễm công thức Excel."""
    cell.value = value
    cell.data_type = "s"
    if align != _PLAIN_ALIGN:
        cell.alignment = align


def _format_date_cell(cell: Any, raw: str, align: Alignment) -> None:
    """Định dạng ô ngày tháng theo chuẩn ISO hoặc chuỗi hiển thị an toàn."""
    clean = (raw or "").strip()
    if not clean:
        _format_cell_text(cell, "", align)
        return

    # Thử bóc tách ngày giờ nhanh qua fromisoformat
    try:
        iso_str = clean[:19].replace(" ", "T")
        d = datetime.fromisoformat(iso_str)
        cell.value = d
        cell.number_format = "yyyy-mm-dd hh:mm" if len(clean) >= 16 else "yyyy-mm-dd"
        if align != _PLAIN_ALIGN:
            cell.alignment = align
        return
    except (ValueError, TypeError):
        pass

    _format_cell_text(cell, clean, align)


def load_user_watchlist(yaml_path: Path) -> dict[str, Any]:
    """Đọc tệp cấu hình Watchlist của người dùng thành tập hợp các từ khóa theo dõi.

    Args:
        yaml_path: Đường dẫn tệp YAML cấu hình của người dùng.

    Returns:
        Từ điển chứa các tập hợp thực thể theo dõi (tickers, industries, entities,...).
    """
    if not yaml_path.exists():
        return {"all_terms": set(), "tickers": set(), "industries": set(), "ticker_regex": None}

    try:
        with open(yaml_path, "r", encoding="utf-8") as f:
            cfg = yaml.safe_load(f) or {}
    except Exception:
        return {"all_terms": set(), "tickers": set(), "industries": set(), "ticker_regex": None}

    tickers = {str(x).strip().upper() for x in (cfg.get("tickers") or []) if str(x).strip()}
    industries = {str(x).strip().upper() for x in (cfg.get("industries") or []) if str(x).strip()}

    all_terms: set[str] = set()
    all_terms.update(tickers)
    all_terms.update(industries)

    for k in ("etfs", "indices", "exchanges", "nations", "themes", "macro", "assets", "institutions", "entities"):
        items = cfg.get(k) or []
        for item in items:
            cleaned = str(item).strip().upper()
            if cleaned:
                all_terms.add(cleaned)

    # Tiền biên dịch biểu thức chính quy cho danh sách mã Ticker hợp lệ (>= 3 ký tự)
    valid_tickers = [t for t in sorted(tickers, key=len, reverse=True) if len(t) >= 3]
    ticker_regex = None
    if valid_tickers:
        ticker_regex = re.compile(r"\b(" + "|".join(re.escape(t) for t in valid_tickers) + r")\b", re.IGNORECASE)

    return {
        "all_terms": all_terms,
        "tickers": tickers,
        "industries": industries,
        "ticker_regex": ticker_regex,
    }


def match_article_watchlist(
    article: dict[str, Any],
    watchlist: dict[str, set[str]],
) -> tuple[bool, str]:
    """Kiểm tra bài viết có khớp với danh mục theo dõi của người dùng hay không.

    Args:
        article: Từ điển dữ liệu bài viết đã chuẩn hóa.
        watchlist: Từ điển các tập hợp từ khóa quan tâm của người dùng.

    Returns:
        Tuple gồm (True nếu khớp / False nếu không, chuỗi các thực thể khớp được).
    """
    all_terms = watchlist.get("all_terms", set())
    if not all_terms:
        return False, ""

    matched_set: set[str] = set()

    # 1. Khớp theo danh sách mã symbols
    raw_symbols = str(article.get("symbols") or "").upper()
    symbols_in_article = {
        s.strip()
        for s in re.split(r"[,;\s]+", raw_symbols)
        if s.strip()
    }
    matched_set.update(symbols_in_article & all_terms)

    # 2. Khớp theo danh sách entities trong bài
    raw_entities = article.get("entities") or []
    for ent in raw_entities:
        ent_str = str(ent).strip().upper()
        if not ent_str:
            continue
        # Tách tiền tố nếu có (vd: TICKER:HPG -> HPG, IND:THEP -> THEP)
        clean_code = ent_str.split(":")[-1].strip()
        if ent_str in all_terms:
            matched_set.add(ent_str)
        if clean_code in all_terms:
            matched_set.add(clean_code)

    # 3. Khớp theo Intent LLM và Intent Code
    for field in ("intent_llm", "intent_code"):
        val = str(article.get(field) or "").upper()
        if val:
            for term in all_terms:
                if len(term) >= 3 and term in val:
                    matched_set.add(term)

    # 4. Khớp theo tiêu đề bài viết cho các mã Ticker độc lập
    title = str(article.get("title") or "")
    if title:
        ticker_rx = watchlist.get("ticker_regex")
        if ticker_rx is not None:
            found = ticker_rx.findall(title)
            if found:
                matched_set.update(f.upper() for f in found)
        else:
            for ticker in watchlist.get("tickers", set()):
                if len(ticker) >= 3:
                    pattern = rf"\b{re.escape(ticker)}\b"
                    if re.search(pattern, title, flags=re.IGNORECASE):
                        matched_set.add(ticker)

    if matched_set:
        sorted_matched = sorted(matched_set)
        return True, ", ".join(sorted_matched)

    return False, ""


def build_delivery_workbook(
    matched_articles: list[dict[str, Any]],
) -> Workbook:
    """Xây dựng bảng tính Excel giao hàng 2 sheets chuẩn mực đơn sắc.

    Args:
        matched_articles: Danh sách các bài viết khớp watchlist kèm trường matched_entities.

    Returns:
        Đối tượng Workbook của openpyxl đã định dạng đầy đủ 2 sheets.
    """
    wb = Workbook()

    # Sheet 1: Tin tức trọng yếu
    ws1 = wb.active
    ws1.title = SHEET_MAIN_NAME

    for col_idx, (_key, label, width, _wrap) in enumerate(SHEET1_COLUMNS, start=1):
        cell = ws1.cell(row=1, column=col_idx)
        _format_cell_text(cell, label, _HEADER_ALIGN)
        cell.font = _HEADER_FONT
        cell.border = _HEADER_BORDER
        ws1.column_dimensions[cell.column_letter].width = width

    for r_idx, row in enumerate(matched_articles, start=2):
        for col_idx, (key, _label, _width, wrap) in enumerate(SHEET1_COLUMNS, start=1):
            cell = ws1.cell(row=r_idx, column=col_idx)
            align = _WRAP_ALIGN if wrap else _PLAIN_ALIGN
            val = row.get(key)

            if key == "date":
                _format_date_cell(cell, str(val or row.get("published_at") or ""), align)
                continue

            if key == "sentiment":
                norm = SENTIMENT_LABELS.get(str(val or "").lower(), str(val or ""))
                _format_cell_text(cell, norm, align)
                continue

            if key == "time_sensitivity":
                norm = TIME_SENSITIVITY_LABELS.get(str(val or "").lower(), str(val or ""))
                _format_cell_text(cell, norm, align)
                continue

            if key == "gold_status":
                norm = GOLD_STATUS_LABELS.get(str(val or "").upper(), str(val or "Full"))
                _format_cell_text(cell, norm, align)
                continue

            if key == "key_points":
                if isinstance(val, (list, tuple)):
                    formatted_kp = "\n".join(f"- {p}" for p in val if str(p).strip())
                else:
                    formatted_kp = str(val or "")
                _format_cell_text(cell, formatted_kp, align)
                continue

            text_val = "" if val is None else str(val)
            _format_cell_text(cell, text_val, align)

            if key == "url" and text_val.startswith(("http://", "https://")):
                cell.hyperlink = text_val

    ws1.freeze_panes = "A2"
    last_col1 = ws1.cell(row=1, column=len(SHEET1_COLUMNS)).column_letter
    ws1.auto_filter.ref = f"A1:{last_col1}{max(ws1.max_row, 1)}"

    # Sheet 2: Luận điểm trích dẫn
    ws2 = wb.create_sheet(title=SHEET_CITATIONS_NAME)

    for col_idx, (_key, label, width, _wrap) in enumerate(SHEET2_COLUMNS, start=1):
        cell = ws2.cell(row=1, column=col_idx)
        _format_cell_text(cell, label, _HEADER_ALIGN)
        cell.font = _HEADER_FONT
        cell.border = _HEADER_BORDER
        ws2.column_dimensions[cell.column_letter].width = width

    sheet2_row = 2
    for article in matched_articles:
        kps = article.get("key_points") or []
        cits = article.get("citations") or []

        if isinstance(kps, str):
            kps = [kps] if kps.strip() else []
        if isinstance(cits, str):
            cits = [cits] if cits.strip() else []

        max_pairs = max(len(kps), len(cits), 1)

        for pair_idx in range(max_pairs):
            kp_text = str(kps[pair_idx]) if pair_idx < len(kps) else ""
            cit_text = str(cits[pair_idx]) if pair_idx < len(cits) else ""

            row_data = {
                "article_id": article.get("article_id", ""),
                "date": article.get("published_at", ""),
                "title": article.get("title", ""),
                "matched_entities": article.get("matched_entities", ""),
                "key_point": kp_text,
                "grounded_citation": cit_text,
                "source_domain": article.get("source_domain", ""),
                "url": article.get("url", ""),
            }

            for col_idx, (col_key, _label, _width, wrap) in enumerate(SHEET2_COLUMNS, start=1):
                cell = ws2.cell(row=sheet2_row, column=col_idx)
                align = _WRAP_ALIGN if wrap else _PLAIN_ALIGN
                val = row_data.get(col_key, "")

                if col_key == "date":
                    _format_date_cell(cell, str(val), align)
                else:
                    _format_cell_text(cell, str(val), align)
                    if col_key == "url" and str(val).startswith(("http://", "https://")):
                        cell.hyperlink = str(val)

            sheet2_row += 1

    ws2.freeze_panes = "A2"
    last_col2 = ws2.cell(row=1, column=len(SHEET2_COLUMNS)).column_letter
    ws2.auto_filter.ref = f"A1:{last_col2}{max(ws2.max_row, 1)}"

    return wb


def distribute_to_users(
    date_str: str,
    parquet_path: str,
    manifest_config_path: str,
    users_config_dir: str,
    output_dir: str = "users/output",
) -> dict[str, Path]:
    """Phân phối các bài viết trong kho Parquet thành các tệp Excel 2 sheets cho từng người dùng.

    Args:
        date_str: Chuỗi ngày xuất bản (vd: '2026-10-08' hoặc '20261008').
        parquet_path: Đường dẫn tệp Parquet tổng hợp đã khử trùng lặp.
        manifest_config_path: Đường dẫn tệp manifest.yaml quản lý bật/tắt người dùng.
        users_config_dir: Thư mục chứa cấu hình YAML danh mục của các người dùng.
        output_dir: Thư mục gốc lưu trữ tệp giao hàng Excel (mặc định 'users/output').

    Returns:
        Từ điển ánh xạ từ tên người dùng sang đường dẫn tệp Excel đã lưu.

    Raises:
        FileNotFoundError: Khi không tìm thấy tệp Parquet tổng hợp đầu vào.
    """
    p_path = Path(parquet_path).resolve()
    if not p_path.exists():
        raise FileNotFoundError(f"Không tìm thấy tệp Parquet tại: {parquet_path}")

    # Đọc cấu hình bật tắt người dùng từ manifest.yaml
    manifest_file = Path(manifest_config_path).resolve()
    enabled_users: set[str] = set()
    global_enabled = True

    if manifest_file.exists():
        try:
            with open(manifest_file, "r", encoding="utf-8") as f:
                manifest_data = yaml.safe_load(f) or {}
            global_enabled = bool(manifest_data.get("enabled", True))
            default_state = bool(manifest_data.get("default", False))
            users_map = manifest_data.get("users", {})

            for user_name, is_on in users_map.items():
                if is_on:
                    enabled_users.add(str(user_name))
        except Exception:
            global_enabled = True

    if not global_enabled:
        return {}

    # Đọc dữ liệu từ tệp Parquet vào danh sách từ điển
    table = pq.read_table(p_path)
    all_articles: list[dict[str, Any]] = table.to_pylist()

    # Thu thập danh sách tệp cấu hình của người dùng
    users_dir = Path(users_config_dir).resolve()
    user_files = list(users_dir.glob("*.yaml")) if users_dir.exists() else []

    results: dict[str, Path] = {}
    out_base = Path(output_dir).resolve()
    clean_date = date_str.replace("/", "-")

    def _export_single_user(u_file: Path) -> tuple[str, Optional[Path]]:
        user_name = u_file.stem
        # Nếu có danh sách người dùng kích hoạt, chỉ xử lý người dùng được bật
        if enabled_users and user_name not in enabled_users:
            return user_name, None

        watchlist = load_user_watchlist(u_file)
        user_matched: list[dict[str, Any]] = []

        for art in all_articles:
            is_matched, matched_entities = match_article_watchlist(art, watchlist)
            if is_matched:
                art_copy = dict(art)
                art_copy["matched_entities"] = matched_entities
                art_copy["date"] = art.get("published_at") or clean_date
                user_matched.append(art_copy)

        # Xây dựng bảng tính Excel 2 sheets
        wb = build_delivery_workbook(user_matched)

        user_target_dir = out_base / user_name
        user_target_dir.mkdir(parents=True, exist_ok=True)
        target_excel_path = user_target_dir / f"{clean_date}.xlsx"

        saved_path, _ = safe_atomic_write(
            target_path=target_excel_path,
            writer_fn_or_content=wb.save,
            binary=True,
            fallback_on_lock=True,
        )
        return user_name, saved_path

    if len(user_files) <= 1:
        for u_file in user_files:
            uname, spath = _export_single_user(u_file)
            if spath is not None:
                results[uname] = spath
    else:
        with concurrent.futures.ThreadPoolExecutor(max_workers=min(len(user_files), 6)) as executor:
            future_to_file = {executor.submit(_export_single_user, uf): uf for uf in user_files}
            for fut in concurrent.futures.as_completed(future_to_file):
                uname, spath = fut.result()
                if spath is not None:
                    results[uname] = spath

    return results
