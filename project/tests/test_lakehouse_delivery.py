"""Kiểm thử phân phối báo cáo người dùng và cấu trúc bảng tính Excel 2 sheets."""

from __future__ import annotations

import tempfile
from pathlib import Path

import openpyxl
import pytest
import yaml

from src.lakehouse.delivery import (
    SHEET_CITATIONS_NAME,
    SHEET_MAIN_NAME,
    build_delivery_workbook,
    distribute_to_users,
    load_user_watchlist,
    match_article_watchlist,
)
from src.lakehouse.ingest import publish_batch
from src.lakehouse.storage import LocalOneDriveStorageAdapter


def test_user_fan_out_delivery_and_two_sheets_structure(tmp_path: Path):
    """Kiểm tra việc lọc chính xác theo Watchlist và định dạng bảng tính Excel 2 sheets.

    Args:
        tmp_path: Thư mục tạm độc lập do pytest cung cấp.
    """
    storage_dir = tmp_path / "news-data"
    users_cfg_dir = tmp_path / "users_cfg"
    manifest_path = tmp_path / "manifest.yaml"
    output_dir = tmp_path / "output"

    storage = LocalOneDriveStorageAdapter(storage_dir)
    users_cfg_dir.mkdir(parents=True)
    date_str = "2026-10-08"

    # 1. Tạo cấu hình người dùng
    anpt_cfg = {"tickers": ["HPG"], "industries": ["THEP"]}
    with open(users_cfg_dir / "AnPT.yaml", "w", encoding="utf-8") as f:
        yaml.safe_dump(anpt_cfg, f)

    phohg_cfg = {"tickers": ["VCB"], "industries": ["NGAN_HANG"]}
    with open(users_cfg_dir / "PhoHG.yaml", "w", encoding="utf-8") as f:
        yaml.safe_dump(phohg_cfg, f)

    tat_cfg = {"tickers": ["HPG", "VCB"]}
    with open(users_cfg_dir / "UserBiTat.yaml", "w", encoding="utf-8") as f:
        yaml.safe_dump(tat_cfg, f)

    manifest_content = {
        "enabled": True,
        "default": False,
        "users": {
            "AnPT": True,
            "PhoHG": True,
            "UserBiTat": False,
        },
    }
    with open(manifest_path, "w", encoding="utf-8") as f:
        yaml.safe_dump(manifest_content, f)

    # 2. Tạo tập dữ liệu bài viết
    records = [
        {
            "article_id": "ART_HPG_1",
            "title": "Hòa Phát đạt lợi nhuận kỷ lục trong quý 3",
            "symbols": "HPG",
            "entities": ["TICKER:HPG", "IND:THEP"],
            "summary": "Tập đoàn Hòa Phát ghi nhận tăng trưởng.",
            "key_points": ["Sản lượng thép tăng 25%", "Biên lợi nhuận cải thiện"],
            "implication": "Thị giá HPG có tiềm năng tăng trưởng ngắn hạn.",
            "sentiment": "pos",
            "time_sensitivity": "today",
            "citations": ["Trích dẫn 1: Doanh thu quý 3 cao", "Trích dẫn 2: Dung Quất chạy 100%"],
            "published_at": "2026-10-08T08:30:00Z",
            "url": "https://fpts.com.vn/news/hpg-q3",
        },
        {
            "article_id": "ART_VCB_1",
            "title": "Vietcombank giữ vững vị thế quán quân lợi nhuận",
            "symbols": "VCB",
            "entities": ["TICKER:VCB"],
            "summary": "Ngân hàng VCB duy trì chất lượng tài sản.",
            "key_points": ["Tỷ lệ nợ xấu ở mức thấp"],
            "implication": "Cổ phiếu VCB là nền tảng phòng thủ.",
            "sentiment": "pos",
            "time_sensitivity": "week",
            "citations": ["Trích dẫn 1: Tỷ lệ bao phủ nợ xấu cao"],
            "published_at": "2026-10-08T09:00:00Z",
            "url": "https://fpts.com.vn/news/vcb-perf",
        },
        {
            "article_id": "ART_OTHER_1",
            "title": "Thị trường bất động sản ghi nhận một số diễn biến mới",
            "symbols": "NVL",
            "entities": ["TICKER:NVL"],
            "summary": "Thị trường chung chưa có nhiều đột phá.",
            "key_points": ["Giao dịch chững lại"],
            "implication": "Cần thêm thời gian để dòng tiền phục hồi.",
            "sentiment": "neu",
            "time_sensitivity": "month",
            "citations": ["Trích dẫn: Giao dịch phân khúc đất nền giảm"],
            "published_at": "2026-10-08T09:30:00Z",
            "url": "https://fpts.com.vn/news/nvl-market",
        },
    ]

    parquet_rel = publish_batch(records, date_str, "dev_test", "b1", storage)
    parquet_abs = storage.get_local_path(parquet_rel)
    assert parquet_abs is not None

    deliveries = distribute_to_users(
        date_str=date_str,
        parquet_path=str(parquet_abs),
        manifest_config_path=str(manifest_path),
        users_config_dir=str(users_cfg_dir),
        output_dir=str(output_dir),
    )

    assert "AnPT" in deliveries
    assert "PhoHG" in deliveries
    assert "UserBiTat" not in deliveries

    # Kiểm tra workbook của AnPT
    anpt_file = deliveries["AnPT"]
    assert anpt_file.exists()
    wb_anpt = openpyxl.load_workbook(anpt_file)
    assert SHEET_MAIN_NAME in wb_anpt.sheetnames
    assert SHEET_CITATIONS_NAME in wb_anpt.sheetnames

    ws1 = wb_anpt[SHEET_MAIN_NAME]
    assert ws1.max_row == 2
    assert ws1.cell(row=2, column=6).value == "Hòa Phát đạt lợi nhuận kỷ lục trong quý 3"
    assert "HPG" in str(ws1.cell(row=2, column=2).value)

    ws2 = wb_anpt[SHEET_CITATIONS_NAME]
    assert ws2.max_row == 3
    assert ws2.cell(row=2, column=5).value == "Sản lượng thép tăng 25%"
    assert ws2.cell(row=2, column=6).value == "Trích dẫn 1: Doanh thu quý 3 cao"
    assert ws2.cell(row=3, column=5).value == "Biên lợi nhuận cải thiện"
    assert ws2.cell(row=3, column=6).value == "Trích dẫn 2: Dung Quất chạy 100%"

    # Kiểm tra workbook của PhoHG
    phohg_file = deliveries["PhoHG"]
    assert phohg_file.exists()
    wb_phohg = openpyxl.load_workbook(phohg_file)
    ws1_pho = wb_phohg[SHEET_MAIN_NAME]
    assert ws1_pho.max_row == 2
    assert ws1_pho.cell(row=2, column=6).value == "Vietcombank giữ vững vị thế quán quân lợi nhuận"


def test_delivery_formatting_security_and_styles(tmp_path: Path):
    """Kiểm tra an toàn chống injection công thức, font tiêu đề và đóng băng dòng.

    Args:
        tmp_path: Thư mục tạm độc lập do pytest cung cấp.
    """
    sample_article = {
        "article_id": "SEC_001",
        "title": "=cmd|' /C calc'!A0",
        "symbols": "HPG",
        "matched_entities": "HPG",
        "summary": "@SUM(1+1)",
        "key_points": ["- Luận điểm 1"],
        "implication": "+12345",
        "citations": ["Trích dẫn an toàn"],
        "url": "https://example.com/test",
        "date": "2026-10-08",
    }

    wb = build_delivery_workbook([sample_article])
    ws1 = wb[SHEET_MAIN_NAME]
    ws2 = wb[SHEET_CITATIONS_NAME]

    # Kiểm tra đóng băng tiêu đề
    assert ws1.freeze_panes == "A2"
    assert ws2.freeze_panes == "A2"

    # Kiểm tra font tiêu đề in đậm
    assert ws1.cell(row=1, column=1).font.bold is True
    assert ws2.cell(row=1, column=1).font.bold is True

    # Kiểm tra chống injection: ô tiêu đề chứa dấu = phải có data_type là chuỗi 's'
    title_cell = ws1.cell(row=2, column=6)
    assert title_cell.data_type == "s"
    assert title_cell.value == "=cmd|' /C calc'!A0"

    # Kiểm tra summary chứa ký tự @ phải là chuỗi 's'
    summary_cell = ws1.cell(row=2, column=11)
    assert summary_cell.data_type == "s"


def test_delivery_global_enabled_false(tmp_path: Path):
    """Kiểm tra trường hợp cờ toàn cục enabled là False trong manifest.yaml.

    Args:
        tmp_path: Thư mục tạm độc lập do pytest cung cấp.
    """
    storage = LocalOneDriveStorageAdapter(tmp_path / "data")
    manifest_p = tmp_path / "manifest.yaml"
    users_dir = tmp_path / "users"
    users_dir.mkdir()

    with open(users_dir / "AnPT.yaml", "w", encoding="utf-8") as f:
        yaml.safe_dump({"tickers": ["HPG"]}, f)

    with open(manifest_p, "w", encoding="utf-8") as f:
        yaml.safe_dump({"enabled": False, "users": {"AnPT": True}}, f)

    rel_p = publish_batch(
        [{"article_id": "1", "title": "Tin HPG", "symbols": "HPG"}],
        "2026-10-08",
        "d1",
        "b1",
        storage,
    )
    abs_p = storage.get_local_path(rel_p)
    assert abs_p is not None

    res = distribute_to_users(
        date_str="2026-10-08",
        parquet_path=str(abs_p),
        manifest_config_path=str(manifest_p),
        users_config_dir=str(users_dir),
        output_dir=str(tmp_path / "out"),
    )
    assert res == {}


def test_delivery_zero_matching_articles(tmp_path: Path):
    """Kiểm tra xử lý khi bài viết không khớp bất kỳ thực thể nào trong watchlist.

    Args:
        tmp_path: Thư mục tạm độc lập do pytest cung cấp.
    """
    storage = LocalOneDriveStorageAdapter(tmp_path / "data")
    manifest_p = tmp_path / "manifest.yaml"
    users_dir = tmp_path / "users"
    users_dir.mkdir()

    with open(users_dir / "AnPT.yaml", "w", encoding="utf-8") as f:
        yaml.safe_dump({"tickers": ["FPT"]}, f)

    with open(manifest_p, "w", encoding="utf-8") as f:
        yaml.safe_dump({"enabled": True, "users": {"AnPT": True}}, f)

    rel_p = publish_batch(
        [{"article_id": "1", "title": "Tin HPG chỉ về thép", "symbols": "HPG"}],
        "2026-10-08",
        "d1",
        "b1",
        storage,
    )
    abs_p = storage.get_local_path(rel_p)
    assert abs_p is not None

    res = distribute_to_users(
        date_str="2026-10-08",
        parquet_path=str(abs_p),
        manifest_config_path=str(manifest_p),
        users_config_dir=str(users_dir),
        output_dir=str(tmp_path / "out"),
    )
    assert "AnPT" in res
    out_wb = openpyxl.load_workbook(res["AnPT"])
    ws1 = out_wb[SHEET_MAIN_NAME]
    # Chỉ có dòng header, 0 dòng dữ liệu
    assert ws1.max_row == 1


def test_distribute_missing_parquet_raises_error(tmp_path: Path):
    """Kiểm tra ném FileNotFoundError khi tệp Parquet đầu vào không tồn tại.

    Args:
        tmp_path: Thư mục tạm độc lập do pytest cung cấp.
    """
    with pytest.raises(FileNotFoundError, match="Không tìm thấy tệp Parquet"):
        distribute_to_users(
            date_str="2026-10-08",
            parquet_path=str(tmp_path / "non_existent.parquet"),
            manifest_config_path=str(tmp_path / "manifest.yaml"),
            users_config_dir=str(tmp_path / "users"),
            output_dir=str(tmp_path / "out"),
        )
