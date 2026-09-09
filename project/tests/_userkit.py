"""Helpers dùng chung cho test lớp NGƯỜI DÙNG: registry + seed article/L1/agent."""
from __future__ import annotations

import json

from src.agent.entities import EntityRegistry
from src.db.store import ArticleStore

ENTITIES = [
    {"entity_id": "TICKER:HPG", "type": "TICKER", "code": "HPG",
     "canonical_name": "Hòa Phát", "aliases": ["Hòa Phát"], "attributes": {}, "sources": []},
    {"entity_id": "TICKER:VCB", "type": "TICKER", "code": "VCB",
     "canonical_name": "Vietcombank", "aliases": ["Vietcombank"], "attributes": {}, "sources": []},
    {"entity_id": "TICKER:VNM", "type": "TICKER", "code": "VNM",
     "canonical_name": "Vinamilk", "aliases": ["Vinamilk"], "attributes": {}, "sources": []},
    {"entity_id": "TICKER:FPT", "type": "TICKER", "code": "FPT",
     "canonical_name": "FPT", "aliases": ["FPT Corp"], "attributes": {}, "sources": []},
    # Entity DIỆN RỘNG — dùng cho test noise filter (nhánh broad_types).
    {"entity_id": "ASSET_CLASS:VANG", "type": "ASSET_CLASS", "code": "VANG",
     "canonical_name": "Vàng", "aliases": ["vàng"], "attributes": {}, "sources": []},
]


def read_delivery(path):
    """Đọc deliverable .xlsx → list[dict] khoá theo FIELD KEY nội bộ (không phải nhãn VN).

    Giữ assertion trong test đọc được (`r["article_id"]`) mà không phụ thuộc câu chữ nhãn.
    """
    from openpyxl import load_workbook

    from src.export.xlsx_delivery import DELIVERY_FIELDS

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(wb.active.iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        return []
    labels = ["" if c is None else str(c) for c in rows[0]]
    key_by_label = {label: key for key, label, _w, _wrap in DELIVERY_FIELDS}
    out = []
    for r in rows[1:]:
        d = {}
        for i, lab in enumerate(labels):
            v = r[i] if i < len(r) else None
            if v is None:
                v = ""
            elif hasattr(v, "strftime"):
                v = v.strftime("%Y-%m-%d")
            d[key_by_label.get(lab, lab)] = str(v)
        out.append(d)
    return out


def delivery_labels(path) -> list[str]:
    """Dòng header (nhãn hiển thị) của deliverable .xlsx."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        first = next(wb.active.iter_rows(values_only=True))
    finally:
        wb.close()
    return ["" if c is None else str(c) for c in first]


def make_registry(subs=None) -> EntityRegistry:
    return EntityRegistry(ENTITIES, subs or {})


def make_store(tmp_path) -> ArticleStore:
    return ArticleStore(db_path=str(tmp_path / "t.db"))


def seed_article(store, aid, *, title="Tin", domain="cafef.vn",
                 published="2026-08-18T09:00:00+07:00") -> None:
    conn = store.connect()
    conn.execute(
        "INSERT OR REPLACE INTO articles "
        "(url, url_title_hash, title, source_domain, published_at, fetched_at) "
        "VALUES (?,?,?,?,?,?)",
        (f"http://x/{aid}", aid, title, domain, published, published))
    conn.commit()
    conn.close()


def seed_l1(store, aid, entity_ids, *, dod_pass=1, title="Tin", etype="TICKER") -> None:
    out = {
        "l1_output_version": "1.0", "article_id": aid, "title": title, "recognized": True,
        "entities": [{"surface": "x", "entity_id": e, "type": etype,
                      "method": "alias", "in_list": True, "confidence": 0.9} for e in entity_ids],
        "categories": {"ticker_company": "done", "etf_fund": "none", "index": "none",
                       "exchange": "none", "industry_sector": "none"},
        "citations": [{"source_span": title}], "confidence": 0.9,
        "processing_metadata": {"agent_provider": "x", "model_used": "m", "timestamp": "t"},
    }
    store.insert_l1_output({
        "article_id": aid, "output_json": json.dumps(out, ensure_ascii=False),
        "recognized": 1, "agent_provider": "x", "model_used": "m", "confidence": 0.9,
        "dod_pass": dod_pass, "dod_reasons": "[]", "created_at": "t"})


def seed_agent(store, aid, *, dod_pass=1, event_type="macro", with_optional=True,
               raw_sha256=None, summary=None, materiality=0.6,
               time_sensitivity="this_week") -> None:
    out = {
        "output_schema_version": "1.0", "article_id": aid,
        "summary": {"abstractive": summary or ("Tóm tắt " + aid), "key_points": ["kp1", "kp2"]},
        "implication": {"text": "Hàm ý", "affected_parties": ["x"], "impact_area": "market"},
        "materiality": {"score": materiality, "time_sensitivity": time_sensitivity},
        "confidence": 0.8,
        "citations": [{"claim": "c", "source_span": "span dài hơn hai mươi ký tự", "source_offset": 0}],
        "processing_metadata": {"agent_provider": "p", "model_used": "m", "timestamp": "t"},
    }
    if with_optional:
        out["sentiment"] = {"overall": -0.3, "polarity": "negative"}
        out["event_type"] = event_type
    store.insert_agent_output({
        "article_id": aid, "raw_sha256": raw_sha256 or ("h_" + aid), "work_item_id": None,
        "output_json": json.dumps(out, ensure_ascii=False), "agent_provider": "p",
        "model_used": "m", "confidence": 0.8, "dod_pass": dod_pass,
        "dod_reasons": "[]", "created_at": "t"})
