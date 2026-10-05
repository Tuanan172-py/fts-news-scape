"""Test dựng bảng tín hiệu insight (chú ý, sentiment, liên kết, hồ sơ nguồn) và sheet Radar chú ý."""

from __future__ import annotations

import json
import math
import statistics
from datetime import date, datetime, timedelta, timezone

import _userkit as k
import pytest
from openpyxl import load_workbook

from src.analytics import signals as sg
from src.export.radar_sheet import SHEET_NAME
from src.export.user_output import UserOutputWriter

VN = timezone(timedelta(hours=7))
SPIKE = date(2026, 9, 25)                       # thứ Sáu
BASE_DAYS = [date(2026, 9, 14), date(2026, 9, 15), date(2026, 9, 16), date(2026, 9, 17),
             date(2026, 9, 18), date(2026, 9, 21), date(2026, 9, 22), date(2026, 9, 23),
             date(2026, 9, 24)]


def at(d: date, hh: int, mm: int = 0) -> str:
    return datetime(d.year, d.month, d.day, hh, mm, tzinfo=VN).isoformat(timespec="seconds")


def seed(store, aid, *, source="cafef.vn", ts, entities=("TICKER:HPG",), sentiment=None,
         cluster=None, role=None, l1=True, l1_source="agent", evidence=None):
    conn = store.connect()
    conn.execute("INSERT OR REPLACE INTO articles (url, url_title_hash, title, source_domain, "
                 "published_at, fetched_at) VALUES (?,?,?,?,?,?)",
                 (f"http://x/{aid}", aid, "Tin", source, ts, ts))
    if cluster:
        conn.execute("INSERT OR REPLACE INTO cluster_members (article_id, cluster_id, role, method, "
                     "evidence) VALUES (?,?,?,?,?)",
                     (aid, cluster, role, "sha", json.dumps(evidence or {})))
    conn.commit()
    conn.close()
    if not l1:
        return
    out = {"entities": [{"entity_id": e, "in_list": True} for e in entities]}
    store.insert_l1_output({"article_id": aid, "output_json": json.dumps(out), "recognized": 1,
                            "agent_provider": "x", "model_used": "m", "confidence": 0.9,
                            "dod_pass": 1, "dod_reasons": "[]", "created_at": ts,
                            "l1_source": l1_source})
    if sentiment is not None:
        store.insert_agent_output({"article_id": aid, "raw_sha256": "h" + aid, "work_item_id": None,
                                   "output_json": json.dumps({"sentiment": sentiment}),
                                   "agent_provider": "p", "model_used": "m", "confidence": 0.8,
                                   "dod_pass": 1, "dod_reasons": "[]", "created_at": ts})


def seed_scenario(store):
    for i, d in enumerate(BASE_DAYS):
        seed(store, f"b{i}", source="base.vn", ts=at(d, 10), sentiment="neutral")
    seed(store, "a1", source="s1.vn", ts=at(SPIKE, 9, 0), entities=("TICKER:HPG", "IND_GICS3:THEP"),
         sentiment="positive", cluster="C", role="canonical")
    seed(store, "a2", source="s2.vn", ts=at(SPIKE, 9, 10), entities=("TICKER:HPG", "IND_GICS3:THEP"),
         sentiment="positive", cluster="C", role="copy")
    seed(store, "a3", source="s3.vn", ts=at(SPIKE, 9, 20), entities=("TICKER:HPG", "IND_GICS3:THEP"),
         sentiment="negative", cluster="C", role="candidate")
    seed(store, "a4", source="s4.vn", ts=at(SPIKE, 9, 40), entities=("TICKER:HPG", "TICKER:VCB"),
         sentiment="neutral")
    seed(store, "a5", source="s4.vn", ts=at(SPIKE, 9, 50), entities=("TICKER:HPG", "TICKER:VCB"))
    seed(store, "a6", source="s1.vn", ts=at(SPIKE, 10, 0), sentiment="positive")
    for i in range(4):                                   # bài đăng cùng phiên nhưng chưa phân tích
        seed(store, f"u{i}", source="s9.vn", ts=at(SPIKE, 11, i), l1=False)


@pytest.fixture
def store(tmp_path):
    return k.make_store(tmp_path)


def build(store, **kw):
    conn = store._connect()
    try:
        return sg.build_all(conn, today=SPIKE, days=30, **kw)
    finally:
        conn.close()


def cell(result, entity, day):
    return next(r for r in result["signal_daily"]
                if r["entity_id"] == entity and r["trade_date"] == day.isoformat())


# -- hàm thuần ----------------------------------------------------------------
@pytest.mark.parametrize("ts,expected", [
    (datetime(2026, 9, 25, 15, 0, tzinfo=VN), date(2026, 9, 28)),    # thứ Sáu sau giờ đóng cửa
    (datetime(2026, 9, 26, 10, 0, tzinfo=VN), date(2026, 9, 28)),    # thứ Bảy
    (datetime(2026, 9, 27, 23, 0, tzinfo=VN), date(2026, 9, 28)),    # Chủ nhật
    (datetime(2026, 9, 28, 14, 45, tzinfo=VN), date(2026, 9, 28)),   # đúng giờ đóng cửa vẫn trong phiên
    (datetime(2026, 9, 28, 14, 46, tzinfo=VN), date(2026, 9, 29)),
    (datetime(2026, 9, 25, 14, 45, tzinfo=VN), date(2026, 9, 25)),
])
def test_trade_date_mapping(ts, expected):
    assert sg.trade_date_of(ts) == expected


def test_sentiment_score_handles_both_output_shapes():
    assert sg.sentiment_score(json.dumps({"sentiment": "positive"})) == 1
    assert sg.sentiment_score(json.dumps({"sentiment": "negative"})) == -1
    assert sg.sentiment_score(json.dumps({"sentiment": {"overall": -0.3, "polarity": "neutral"}})) == 0
    assert sg.sentiment_score(json.dumps({"sentiment": "lạ"})) is None
    assert sg.sentiment_score(None) is None and sg.sentiment_score("không phải json") is None


def test_entity_ids_keeps_only_allowed_prefixes():
    raw = json.dumps({"entities": [{"entity_id": "TICKER:HPG"}, {"entity_id": "MACRO_FX"},
                                   {"entity_id": "EXCHANGE:HOSE"}, {"entity_id": "ASSET_CLASS:VANG"},
                                   {"entity_id": "IND_GICS3:THEP"}, {"surface": "x"}]})
    assert sg.entity_ids(raw) == {"TICKER:HPG", "MACRO_FX", "ASSET_CLASS:VANG", "IND_GICS3:THEP"}
    assert sg.entity_ids(raw, sg.LINK_PREFIXES) == {"TICKER:HPG", "IND_GICS3:THEP"}


def test_zscore_uses_sd_floor():
    assert sg.zscore(6, [1, 1, 1, 1, 1]) == pytest.approx(10.0)       # sd = 0 nên dùng sàn 0,5
    assert sg.zscore(3, [1, 3, 5]) == pytest.approx(0.0)


# -- signal_daily -------------------------------------------------------------
def test_attention_metrics_on_spike_day(store):
    seed_scenario(store)
    r = cell(build(store), "TICKER:HPG", SPIKE)
    assert (r["n_articles"], r["n_sources"], r["n_stories"]) == (6, 4, 4)
    assert r["share"] == pytest.approx(1.0)                  # mọi bài phân tích trong phiên đều có HPG
    assert r["ama_z"] == pytest.approx(10.0)                 # nền 1 bài mỗi phiên, sàn sd 0,5
    assert r["first_source"] == "s1.vn"
    assert r["n_updates"] == 1
    assert r["stale_ratio"] == pytest.approx(1 / 6)          # chỉ bài copy, candidate chưa có novelty
    assert r["cascade_minutes"] == pytest.approx(20.0)       # bài thứ 3 của cụm C sau bài đầu 20 phút
    assert r["coverage_pct"] == pytest.approx(60.0)          # 6 bài phân tích trên 10 bài đăng cùng phiên


def test_sentiment_story_volume_and_dispersion(store):
    seed_scenario(store)
    res = build(store)
    r = cell(res, "TICKER:HPG", SPIKE)
    assert r["net_sent_volume"] == pytest.approx((1 + 1 - 1 + 0 + 1) / 5)
    assert r["net_sent_story"] == pytest.approx((1 + 0 + 1) / 3)      # cụm C một phiếu từ bài gốc
    means = [1, 1, -1, 0]                                             # s1 (a1, a6), s2, s3, s4 (a4)
    assert r["dispersion"] == pytest.approx(statistics.pstdev(means))
    base = {p["source_domain"]: (p["pos_rate"] or 0) - (p["neg_rate"] or 0) for p in res["source_profile"]}
    scores = {"a1": ("s1.vn", 1), "a2": ("s2.vn", 1), "a3": ("s3.vn", -1), "a4": ("s4.vn", 0),
              "a6": ("s1.vn", 1)}
    assert r["net_sent_norm"] == pytest.approx(
        statistics.fmean(s - base[src] for src, s in scores.values()))


def test_baseline_needs_five_sessions(store):
    seed_scenario(store)
    res = build(store)
    assert cell(res, "TICKER:HPG", BASE_DAYS[0])["ama_z"] is None
    assert cell(res, "TICKER:HPG", BASE_DAYS[4])["ama_z"] is None     # mới có 4 phiên trước
    assert cell(res, "TICKER:HPG", BASE_DAYS[5])["ama_z"] is not None


def test_sent_shift_is_difference_from_previous_sessions(store):
    seed_scenario(store)
    res = build(store)
    last_base = cell(res, "TICKER:HPG", BASE_DAYS[-1])
    assert last_base["sent_shift"] == pytest.approx(0.0)              # nền toàn trung tính, cùng mức nguồn
    spike = cell(res, "TICKER:HPG", SPIKE)
    prior = [cell(res, "TICKER:HPG", d)["net_sent_norm"] for d in BASE_DAYS[-5:]]
    assert spike["sent_shift"] == pytest.approx(spike["net_sent_norm"] - statistics.fmean(prior))


def test_missing_cluster_rows_make_each_article_its_own_story(store):
    for i in range(3):
        seed(store, f"x{i}", source=f"s{i}.vn", ts=at(SPIKE, 9, i), sentiment="positive")
    r = cell(build(store), "TICKER:HPG", SPIKE)
    assert r["n_stories"] == 3 and r["stale_ratio"] == 0 and r["cascade_minutes"] is None
    assert r["n_updates"] == 0


def test_candidate_with_novelty_none_counts_as_stale(store):
    seed(store, "c1", ts=at(SPIKE, 9, 0), cluster="K", role="canonical")
    seed(store, "c2", source="s2.vn", ts=at(SPIKE, 9, 5), cluster="K", role="candidate",
         evidence={"novelty": "none"})
    seed(store, "c3", source="s3.vn", ts=at(SPIKE, 9, 6), cluster="K", role="candidate",
         evidence={"novelty": "new_facts"})
    r = cell(build(store), "TICKER:HPG", SPIKE)
    assert r["stale_ratio"] == pytest.approx(1 / 3) and r["n_updates"] == 2


def test_code_first_and_failed_results_are_excluded(store):
    seed(store, "ok", ts=at(SPIKE, 9, 0))
    seed(store, "cf", ts=at(SPIKE, 9, 1), l1_source="code_first")
    seed(store, "inh", source="s2.vn", ts=at(SPIKE, 9, 2), l1_source="inherited")
    r = cell(build(store), "TICKER:HPG", SPIKE)
    assert r["n_articles"] == 2                                       # code_first bị loại, inherited được tính


# -- liên kết và hồ sơ nguồn --------------------------------------------------
def test_entity_links_threshold_and_pmi(store):
    seed_scenario(store)
    links = build(store)["entity_links"]
    pairs = {(x["entity_a"], x["entity_b"]): x for x in links}
    assert ("IND_GICS3:THEP", "TICKER:HPG") in pairs
    assert ("TICKER:HPG", "TICKER:VCB") not in pairs                  # chỉ 2 lần đồng xuất hiện
    link = pairs[("IND_GICS3:THEP", "TICKER:HPG")]
    n_total, n_hpg, n_thep = 15, 15, 3
    assert link["n_co"] == 3
    assert link["pmi"] == pytest.approx(math.log2(3 * n_total / (n_hpg * n_thep)))
    assert all(x["entity_a"] < x["entity_b"] for x in links)


def test_source_profile_lead_and_original_rates(store):
    seed_scenario(store)
    prof = {p["source_domain"]: p for p in build(store)["source_profile"]}
    assert prof["s1.vn"]["lead_rate"] == pytest.approx(1.0)           # a1 dẫn đầu cụm C ba nguồn
    assert prof["s2.vn"]["lead_rate"] == 0.0 and prof["s3.vn"]["lead_rate"] == 0.0
    assert prof["s4.vn"]["lead_rate"] is None                         # không có bài trong cụm đa nguồn
    assert prof["s1.vn"]["original_rate"] == 1.0
    assert prof["s2.vn"]["original_rate"] == 0.0                      # a2 là copy
    assert prof["s1.vn"]["pos_rate"] == 1.0 and prof["s3.vn"]["neg_rate"] == 1.0


# -- ghi DB -------------------------------------------------------------------
def _snapshot(store):
    conn = store._connect()
    try:
        return {
            "daily": conn.execute("SELECT entity_id, trade_date, n_articles, ama_z, net_sent_norm "
                                  "FROM signal_daily ORDER BY 1, 2").fetchall(),
            "links": conn.execute("SELECT * FROM entity_links ORDER BY 1, 2").fetchall(),
            "prof": conn.execute("SELECT * FROM source_profile ORDER BY 1").fetchall()}
    finally:
        conn.close()


def test_build_is_idempotent_and_dry_run_writes_nothing(store):
    seed_scenario(store)
    build(store, write=False)
    assert _snapshot(store) == {"daily": [], "links": [], "prof": []}
    build(store)
    first = _snapshot(store)
    assert first["daily"] and first["links"] and first["prof"]
    build(store)
    assert _snapshot(store) == first


def test_rebuild_drops_rows_that_no_longer_qualify(store):
    seed_scenario(store)
    build(store)
    conn = store._connect()
    conn.execute("DELETE FROM l1_outputs WHERE article_id='a6'")
    conn.commit()
    conn.close()
    build(store)
    conn = store._connect()
    n = conn.execute("SELECT n_articles FROM signal_daily WHERE entity_id='TICKER:HPG' "
                     "AND trade_date=?", (SPIKE.isoformat(),)).fetchone()[0]
    conn.close()
    assert n == 5


# -- sheet Radar chú ý --------------------------------------------------------
DAY = "2026-08-18"


def _deliver(tmp_path, build_signals: bool):
    store = k.make_store(tmp_path)
    reg = k.make_registry({"AnPT": {"TICKER:HPG"}})
    k.seed_article(store, "a1")
    k.seed_l1(store, "a1", ["TICKER:HPG"])
    k.seed_agent(store, "a1")
    k.seed_article(store, "a2", domain="vietstock.vn")
    k.seed_l1(store, "a2", ["TICKER:VNM"])
    if build_signals:
        conn = store._connect()
        sg.build_all(conn, today=date(2026, 8, 18), days=30)
        conn.close()
    UserOutputWriter(store, reg, output_root=tmp_path / "out").write(date=DAY)
    return tmp_path / "out" / "AnPT" / f"{DAY}.xlsx"


def test_radar_sheet_added_and_filtered_by_watchlist(tmp_path):
    path = _deliver(tmp_path, True)
    wb = load_workbook(path, read_only=True)
    try:
        assert wb.sheetnames[0] == "Watchlist News" and wb.sheetnames[1] == SHEET_NAME
        rows = list(wb[SHEET_NAME].iter_rows(values_only=True))
    finally:
        wb.close()
    assert rows[0][:2] == ("Phiên", "Thực thể")
    assert [r[1] for r in rows[1:]] == ["TICKER:HPG"]                 # VNM không thuộc watchlist
    assert rows[1][2] == 1


def test_radar_sheet_omitted_when_no_signals(tmp_path):
    path = _deliver(tmp_path, False)
    wb = load_workbook(path, read_only=True)
    try:
        assert wb.sheetnames == ["Watchlist News"]
    finally:
        wb.close()
